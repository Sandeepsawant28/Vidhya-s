import socket
from datetime import datetime
from decimal import Decimal
from django.utils import timezone
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .models import AuditLog, Bill

def get_client_ip(request):
    """Extract client IP from request headers or remote address."""
    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        return x_forwarded_for.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR', '')

def get_local_ip():
    """Detects the PC's primary LAN IP address (e.g. 192.168.1.15)."""
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        # Doesn't need to be reachable, just triggers local interface selection
        s.connect(('10.255.255.255', 1))
        ip = s.getsockname()[0]
    except Exception:
        try:
            ip = socket.gethostbyname(socket.gethostname())
        except Exception:
            ip = '127.0.0.1'
    finally:
        s.close()
    return ip

def log_audit(request, action, details):
    """Creates a centralized audit log entry."""
    user = request.user if request and request.user.is_authenticated else None
    ip = get_client_ip(request) if request else ''
    AuditLog.objects.create(
        user=user,
        action=action,
        details=details,
        ip_address=ip
    )

def generate_bill_number():
    """Generates unique sequential bill number: BILL-YYYYMMDD-XXXX."""
    today = timezone.localtime().strftime('%Y%m%d')
    prefix = f"BILL-{today}-"
    last_bill = Bill.objects.filter(bill_number__startswith=prefix).order_by('-bill_number').first()
    if last_bill:
        try:
            last_seq = int(last_bill.bill_number.split('-')[-1])
            new_seq = last_seq + 1
        except Exception:
            new_seq = 1
    else:
        new_seq = 1
    return f"{prefix}{new_seq:04d}"

def export_queryset_to_excel(filename, sheet_name, headers, data_rows):
    """
    Exports data to a styled Excel sheet using openpyxl and returns an HttpResponse.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:30]

    # Styles
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Write headers
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    # Write data rows
    for row_idx, row in enumerate(data_rows, start=2):
        formatted_row = []
        for val in row:
            if isinstance(val, Decimal):
                formatted_row.append(float(val))
            elif isinstance(val, (datetime, timezone.datetime)):
                formatted_row.append(val.strftime('%Y-%m-%d %H:%M'))
            else:
                formatted_row.append(val)
        ws.append(formatted_row)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal='right')

    # Auto-adjust column width
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
